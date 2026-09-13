#!/usr/bin/env python3
"""Build the final workbook: metadata + embedded thumbnail + timed transcript.

Run after you have the two local folders:

    python3 transcript.py data/links.txt -o raw/       # timed transcripts
    python3 thumbs.py  fetch data/links.txt -o thumbs/ # thumbnail images
    python3 finalize.py -o data/out/outliers_full.xlsx

The thumbnail is embedded as a real picture in the row, not a link. The
transcript goes in the next column as "[m:ss] line" text. Excel caps a cell
at 32,767 characters, so anything longer is written to
data/out/transcripts/<id>.txt and the cell carries the opening plus a
pointer -- the 53-minute entry needs this, the rest fit inline.
"""
import argparse, csv, glob, os, re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

CELL_MAX = 32767
THUMB_W, THUMB_H = 240, 135          # 16:9, keeps rows readable
HEAD = PatternFill("solid", fgColor="1F1F1F")
TS = re.compile(r"(\d{2}):(\d{2}):(\d{2})[,.](\d{3})\s*-->")


def srt_to_timed(path):
    """'[m:ss] spoken line' per cue, blank and duplicate cues dropped."""
    out, t, buf, seen = [], None, [], ""
    def flush():
        nonlocal buf, seen
        txt = re.sub(r"\s+", " ", " ".join(buf)).strip()
        if txt and txt != seen and t is not None:
            m, s = divmod(int(t), 60)
            out.append(f"[{m}:{s:02d}] {txt}")
            seen = txt
        buf = []
    for line in open(path, encoding="utf-8", errors="ignore"):
        m = TS.search(line)
        if m:
            flush()
            g = [int(x) for x in m.groups()]
            t = g[0] * 3600 + g[1] * 60 + g[2]
        elif line.strip() and not line.strip().isdigit():
            buf.append(re.sub(r"<[^>]+>", "", line).strip())
    flush()
    return "\n".join(out)


def find(dirname, vid, exts):
    for e in exts:
        hit = glob.glob(os.path.join(dirname, f"{vid}*.{e}"))
        if hit:
            return sorted(hit)[0]
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--videos", default="data/out/videos.csv")
    ap.add_argument("--thumbs", default="thumbs")
    ap.add_argument("--raw", default="raw")
    ap.add_argument("-o", "--out", default="data/out/outliers_full.xlsx")
    ap.add_argument("--seed", default="data/seed.tsv")
    a = ap.parse_args()

    rows = list(csv.DictReader(open(a.videos, encoding="utf-8")))
    rows.sort(key=lambda r: -int(r["views"]))
    mult = {}
    if os.path.exists(a.seed):
        for r in csv.DictReader(open(a.seed, encoding="utf-8"), delimiter="\t"):
            mult[r["video_id"]] = r["multiplier"]

    tdir = os.path.join(os.path.dirname(a.out), "transcripts")
    os.makedirs(tdir, exist_ok=True)
    cache = os.path.join(os.path.dirname(a.out), "_thumbcache")
    os.makedirs(cache, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Outliers"
    cols = ["#", "Title", "Channel", "Multiplier", "Subs", "Views", "Views/day",
            "Uploaded", "Duration", "Like %", "Thumbnail", "Transcript (timed)"]
    ws.append(cols)
    for i, _ in enumerate(cols, 1):
        c = ws.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = HEAD
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)

    got_t = got_x = 0
    for n, r in enumerate(rows, 1):
        vid = r["video_id"]
        ws.append([n, r["title"], r["channel"], mult.get(vid, ""),
                   int(r["channel_subs"]), int(r["views"]),
                   round(float(r["views_per_day"])), r["published"],
                   r["duration"], float(r["like_rate_pct"] or 0), "", ""])
        row = n + 1
        ws.row_dimensions[row].height = THUMB_H * 0.78

        # --- thumbnail, embedded as a picture
        src = find(a.thumbs, vid, ("jpg", "jpeg", "png", "webp"))
        if src:
            dst = os.path.join(cache, f"{vid}.png")
            im = PILImage.open(src).convert("RGB").resize((THUMB_W, THUMB_H))
            im.save(dst)
            pic = XLImage(dst)
            pic.width, pic.height = THUMB_W, THUMB_H
            ws.add_image(pic, f"K{row}")
            got_x += 1
        else:
            ws.cell(row, 11).value = f"https://i.ytimg.com/vi/{vid}/maxresdefault.jpg"

        # --- timed transcript
        srt = find(a.raw, vid, ("srt", "vtt"))
        if srt:
            text = srt_to_timed(srt)
            side = os.path.join(tdir, f"{vid}.txt")
            open(side, "w", encoding="utf-8").write(text)
            if len(text) > CELL_MAX:
                keep = text[:CELL_MAX - 400].rsplit("\n", 1)[0]
                text = keep + f"\n\n[... truncated at Excel's 32,767 character cell limit — full text in transcripts/{vid}.txt]"
            ws.cell(row, 12).value = text
            got_t += 1
        else:
            ws.cell(row, 12).value = "(no transcript file found in %s/)" % a.raw
        ws.cell(row, 12).alignment = Alignment(vertical="top", wrap_text=False)

    widths = {1: 4, 2: 44, 3: 15, 4: 10, 5: 8, 6: 11, 7: 10, 8: 11, 9: 9, 10: 8,
              11: THUMB_W / 7.0, 12: 90}
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"

    wb.save(a.out)
    print(f"wrote {a.out}")
    print(f"  {got_x}/{len(rows)} thumbnails embedded   (missing ones show the URL instead)")
    print(f"  {got_t}/{len(rows)} transcripts inlined   (full text also in {tdir}/)")
    if got_x < len(rows):
        print(f"\n  run:  python3 thumbs.py fetch data/links.txt -o {a.thumbs}/")
    if got_t < len(rows):
        print(f"  run:  python3 transcript.py data/links.txt -o {a.raw}/")


if __name__ == "__main__":
    main()
